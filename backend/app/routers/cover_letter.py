from __future__ import annotations

import io
import json
from typing import Any, Dict, Literal

from docx import Document
from docx.shared import Pt
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

from app.config import Settings, get_settings
from app.services.openai_client import OpenAIClient, OpenAIError

router = APIRouter(prefix="/api", tags=["cover-letter"])

Tone = Literal["professional", "enthusiastic", "concise", "friendly", "formal"]


class CoverLetterRequest(BaseModel):
    cv_text: str = Field(..., min_length=20)
    jd_text: str = Field(..., min_length=20)
    tone: Tone = "professional"


class CoverLetterResponse(BaseModel):
    letter: str
    tone_used: Tone


_COVER_LETTER_SCHEMA: Dict[str, Any] = {
    "type": "object",
    "description": "A tailored cover letter for the candidate.",
    "properties": {
        "letter": {
            "type": "string",
            "description": (
                "A complete, ready-to-send cover letter of 3-5 short paragraphs. "
                "Plain text, no markdown, no placeholders like [Your Name]."
            ),
        },
        "tone_used": {
            "type": "string",
            "enum": ["professional", "enthusiastic", "concise", "friendly", "formal"],
        },
    },
    "required": ["letter", "tone_used"],
    "additionalProperties": False,
}


def _build_system_prompt(tone: Tone) -> str:
    return (
        "You are an expert cover-letter writer. Craft a tailored cover letter using ONLY facts "
        "that appear in the candidate's CV and the provided job description. Do not invent "
        "employers, dates, or achievements. Keep it 3-5 short paragraphs, strong opening hook, "
        "one paragraph on fit with 2-3 concrete strengths from the CV, one paragraph on "
        "enthusiasm for the role/company, a brief closing. Avoid clichés ('dynamic team player'), "
        "placeholders, and markdown. Respond ONLY with JSON matching the schema. "
        f"Tone: {tone}."
    )


def _get_client(settings: Settings = Depends(get_settings)) -> OpenAIClient:
    return OpenAIClient(settings)


@router.post("/cover-letter", response_model=CoverLetterResponse)
async def generate_cover_letter(
    req: CoverLetterRequest,
    client: OpenAIClient = Depends(_get_client),
) -> CoverLetterResponse:
    user_prompt = (
        f"**CV TEXT:**\n---\n{req.cv_text}\n---\n\n"
        f"**JOB DESCRIPTION:**\n---\n{req.jd_text}\n---"
    )
    payload: Dict[str, Any] = {
        "messages": [
            {"role": "system", "content": _build_system_prompt(req.tone)},
            {"role": "user", "content": user_prompt},
        ],
        "temperature": 0.4,
        "response_format": {
            "type": "json_schema",
            "json_schema": {"name": "cover_letter", "schema": _COVER_LETTER_SCHEMA, "strict": True},
        },
    }

    try:
        completion = await client.chat_completion(payload)
    except OpenAIError as exc:
        raise HTTPException(status_code=exc.status_code, detail=exc.detail) from exc

    try:
        content = completion["choices"][0]["message"]["content"]
        parsed = json.loads(content)
    except (KeyError, IndexError, ValueError, TypeError) as exc:
        raise HTTPException(
            status_code=502,
            detail=f"Malformed OpenAI response: {exc!s}",
        ) from exc

    return CoverLetterResponse(**parsed)


class CoverLetterDocxRequest(BaseModel):
    letter: str = Field(..., min_length=10)
    filename: str = Field(default="cover-letter")


def _build_docx(text: str) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    for paragraph in text.split("\n\n"):
        cleaned = paragraph.strip()
        if not cleaned:
            continue
        p = doc.add_paragraph(cleaned)
        p.paragraph_format.space_after = Pt(8)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


@router.post("/cover-letter/docx")
async def download_cover_letter_docx(req: CoverLetterDocxRequest) -> StreamingResponse:
    raw = _build_docx(req.letter)
    safe_name = "".join(c for c in req.filename if c.isalnum() or c in ("-", "_")) or "cover-letter"
    headers = {"Content-Disposition": f'attachment; filename="{safe_name}.docx"'}
    return StreamingResponse(
        io.BytesIO(raw),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers=headers,
    )
